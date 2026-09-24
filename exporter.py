"""Excel export for cached MCP query results."""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from result_store import ResultStore


class ExportError(Exception):
    pass


def _safe_filename(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("_") or "dsn_extract"


def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or "Sheet"
    candidate = base
    i = 2
    while candidate in used:
        suffix = f"_{i}"
        candidate = (base[: 31 - len(suffix)] + suffix)
        i += 1
    used.add(candidate)
    return candidate


def export_query_to_xlsx(
    store: ResultStore,
    query_id: str,
    export_dir: Path,
) -> Path:
    meta = store.get_metadata(query_id)
    export_dir.mkdir(parents=True, exist_ok=True)

    report = str(meta.get("report", "dsn_query"))
    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    out_path = export_dir / f"{timestamp}-{_safe_filename(report)}-{query_id[:8]}.xlsx"

    wb = Workbook()
    default = wb.active
    wb.remove(default)
    used_names: set[str] = set()
    header_font = Font(bold=True)

    for dataset_name, info in meta.get("datasets", {}).items():
        ws = wb.create_sheet(_sheet_name(dataset_name, used_names))
        columns = list(info.get("columns", []))
        ws.append(columns)
        for cell in ws[1]:
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")

        widths = [len(str(c)) for c in columns]
        row_idx = 1
        for row in store.iter_rows(query_id, dataset_name):
            ws.append(row)
            row_idx += 1
            for i, value in enumerate(row):
                if i < len(widths) and value is not None and row_idx <= 101:
                    widths[i] = max(widths[i], len(str(value)))

        for i, col in enumerate(columns, start=1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = min(50, max(10, widths[i - 1] + 2))
            col_l = str(col).lower()
            if any(k in col_l for k in ("revenue", "amount", "price")):
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, i).number_format = "#,##0.00"
            elif any(k in col_l for k in ("count", "txns", "quantity")):
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, i).number_format = "#,##0"

        ws.freeze_panes = "A2"
        if columns:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, ws.max_row)}"

    info_ws = wb.create_sheet("_pull_info")
    info_ws.append(["field", "value"])
    for cell in info_ws[1]:
        cell.font = header_font

    pull_info: list[tuple[str, Any]] = [
        ("query_id", query_id),
        ("report", meta.get("report")),
        ("business_rule_version", meta.get("business_rule_version")),
        ("generated_utc", datetime.now(timezone.utc).isoformat()),
        ("scope", json.dumps(meta.get("scope", {}), ensure_ascii=False)),
        ("methodology", json.dumps(meta.get("methodology", {}), ensure_ascii=False)),
        ("warnings", json.dumps(meta.get("warnings", []), ensure_ascii=False)),
    ]
    for dataset_name, info in meta.get("datasets", {}).items():
        pull_info.extend(
            [
                (f"dataset.{dataset_name}.rows", info.get("row_count")),
                (f"dataset.{dataset_name}.truncated", info.get("truncated")),
                (f"dataset.{dataset_name}.template", info.get("template_name")),
                (f"dataset.{dataset_name}.query_hash", info.get("query_hash")),
            ]
        )
    for field, value in pull_info:
        info_ws.append([field, value])
    info_ws.column_dimensions["A"].width = 42
    info_ws.column_dimensions["B"].width = 110
    info_ws.freeze_panes = "A2"

    try:
        wb.save(out_path)
    except PermissionError as exc:
        raise ExportError(
            f"Could not write '{out_path}'. If the file already exists and is open "
            "in Excel, close it and retry."
        ) from exc

    store.set_export(query_id, out_path)
    return out_path
